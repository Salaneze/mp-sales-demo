"""Таблица: «Заказы», «Остатки», «По дням». Excel всегда, Google Sheets по ключам."""
import logging
from decimal import Decimal
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from .notify import MP_NAMES

log = logging.getLogger(__name__)


def sheets(db) -> list[tuple[str, list[str], list[list]]]:
    orders = [
        [r["date"].replace("T", " "), MP_NAMES.get(r["mp"], r["mp"]), r["sku"], r["name"],
         float(Decimal(r["price"])), r["warehouse"], "отмена" if r["cancelled"] else ""]
        for r in db.execute("SELECT * FROM orders ORDER BY date DESC")
    ]
    stocks = [
        [MP_NAMES.get(r["mp"], r["mp"]), r["sku"], r["name"], r["warehouse"], r["qty"]]
        for r in db.execute("SELECT * FROM stocks ORDER BY mp, sku, warehouse")
    ]
    days: dict[tuple[str, str], list] = {}
    for r in db.execute("SELECT substr(date, 1, 10) d, mp, price, cancelled FROM orders"):
        row = days.setdefault((r["d"], r["mp"]), [0, Decimal(0), 0])
        if r["cancelled"]:
            row[2] += 1
        else:
            row[0] += 1
            row[1] += Decimal(r["price"])
    by_day = [[d, MP_NAMES.get(mp, mp), n, float(s), c]
              for (d, mp), (n, s, c) in sorted(days.items(), reverse=True)]
    return [
        ("Заказы", ["Дата", "Площадка", "Артикул", "Товар", "Цена", "Склад", "Статус"], orders),
        ("Остатки", ["Площадка", "Артикул", "Товар", "Склад", "Остаток"], stocks),
        ("По дням", ["День", "Площадка", "Заказов", "Выручка", "Отмен"], by_day),
    ]


def to_xlsx(db, path: Path) -> None:
    wb = Workbook()
    wb.remove(wb.active)
    for title, header, rows in sheets(db):
        ws = wb.create_sheet(title)
        ws.append(header)
        for row in rows:
            ws.append(row)
        for cell in ws[1]:
            cell.font = Font(bold=True)
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for i, col in enumerate(header, 1):
            width = max([len(col)] + [len(str(r[i - 1])) for r in rows[:300]])
            ws.column_dimensions[get_column_letter(i)].width = min(width + 2, 50)
    path.parent.mkdir(parents=True, exist_ok=True)
    try:
        wb.save(path)
    except PermissionError:
        # Файл открыт в Excel — Windows держит блокировку
        alt = path.with_stem(path.stem + "_new")
        wb.save(alt)
        log.warning("%s открыт в другой программе, сохранил в %s", path, alt)


def to_gsheet(db, sheet_id: str, credentials: str) -> None:
    import gspread  # нужен только тем, кто выгружает в Google

    book = gspread.service_account(filename=credentials).open_by_key(sheet_id)
    for title, header, rows in sheets(db):
        try:
            ws = book.worksheet(title)
        except gspread.WorksheetNotFound:
            ws = book.add_worksheet(title, rows=len(rows) + 10, cols=len(header))
        ws.clear()
        ws.update([header] + rows, value_input_option="USER_ENTERED")
        ws.freeze(rows=1)
