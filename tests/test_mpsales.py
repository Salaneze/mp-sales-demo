"""python -m unittest"""
import sqlite3
import tempfile
import unittest
from datetime import datetime, timedelta
from decimal import Decimal
from pathlib import Path

from openpyxl import load_workbook

from mpsales import notify, report, store
from mpsales.models import Order, Stock
from mpsales.sources import demo
from mpsales.sources.wb import WB, parse_remains

T0 = datetime(2026, 10, 18, 12, 0)


def db():
    con = sqlite3.connect(":memory:")
    con.row_factory = sqlite3.Row
    con.executescript(store.SCHEMA)
    return con


def order(i, cancelled=False, price="1000", day=T0):
    return Order("wb", f"srid-{i}", day, f"SKU{i}", f"Товар {i}", Decimal(price), "Коледино", cancelled)


class WbParseTest(unittest.TestCase):
    # Строка ответа из документации WB (dev.wildberries.ru, 18.10.2026)
    DOC_ORDER = {
        "date": "2022-03-04T18:08:31", "lastChangeDate": "2022-03-06T10:11:07",
        "warehouseName": "Склад WB РФ", "supplierArticle": "12345", "nmId": 1234567,
        "subject": "Мультистайлеры", "brand": "Тест", "techSize": "0",
        "totalPrice": 1887, "discountPercent": 18, "finishedPrice": 1145, "priceWithDisc": 1547,
        "isCancel": True, "srid": "11.rf9ef11fce1684117b0nhj96222982382.3.0",
    }

    def test_order(self):
        o = WB._order(self.DOC_ORDER)
        self.assertEqual(o.order_id, self.DOC_ORDER["srid"])
        self.assertEqual(o.sku, "12345")
        self.assertEqual(o.name, "Мультистайлеры Тест")  # размер «0» не показываем
        self.assertEqual(o.price, Decimal("1547"))
        self.assertTrue(o.cancelled)

    def test_remains_skip_pseudo_warehouses(self):
        rows = [{"brand": "Wonderful", "subjectName": "Фотоальбомы", "vendorCode": "41058", "nmId": 1,
                 "warehouses": [
                     {"warehouseName": "В пути до получателей", "quantity": 14},
                     {"warehouseName": "В пути возвраты на склад WB", "quantity": 4},
                     {"warehouseName": "Всего находится на складах", "quantity": 267},
                     {"warehouseName": "Склад WB РФ", "quantity": 134},
                     {"warehouseName": "Минск", "quantity": 123},
                 ]}]
        stocks = parse_remains(rows)
        self.assertEqual({s.warehouse for s in stocks}, {"Склад WB РФ", "Минск"})
        self.assertEqual(sum(s.qty for s in stocks), 257)


class StoreTest(unittest.TestCase):
    def test_orders_dedup_and_cancel(self):
        con = db()
        new, canc = store.save_orders(con, [order(1), order(2)])
        self.assertEqual((len(new), len(canc)), (2, 0))
        # WB отдал тот же заказ повторно, уже отменённым
        new, canc = store.save_orders(con, [order(1, cancelled=True), order(2), order(3)])
        self.assertEqual([o.order_id for o in new], ["srid-3"])
        self.assertEqual([o.order_id for o in canc], ["srid-1"])
        self.assertEqual(con.execute("SELECT count(*) FROM orders").fetchone()[0], 3)

    def test_day_summary_exact_money(self):
        con = db()
        store.save_orders(con, [order(1, price="0.10"), order(2, price="0.20"), order(3, cancelled=True)])
        s = store.day_summary(con, T0)
        self.assertEqual(s["orders"], 2)
        self.assertEqual(s["revenue"], Decimal("0.30"))


class DemoTest(unittest.TestCase):
    def test_deterministic_and_window_independent(self):
        a = demo.orders("demo", T0 - timedelta(hours=10), T0)
        b = demo.orders("demo", T0 - timedelta(hours=10), T0 - timedelta(hours=5)) + \
            demo.orders("demo", T0 - timedelta(hours=5), T0)
        self.assertEqual([(o.order_id, o.sku, o.price) for o in a], [(o.order_id, o.sku, o.price) for o in b])

    def test_stock_melts(self):
        full = sum(s.qty for s in demo.stocks("demo", {}) if s.sku == "TEE-M")
        after = sum(s.qty for s in demo.stocks("demo", {"TEE-M": 5}) if s.sku == "TEE-M")
        self.assertEqual(full - after, 5)
        self.assertEqual(sum(s.qty for s in demo.stocks("demo", {"TEE-M": 999}) if s.sku == "TEE-M"), 0)


class ReportTest(unittest.TestCase):
    def test_xlsx_sheets(self):
        con = db()
        store.save_orders(con, [order(1), order(2, day=T0 - timedelta(days=1)), order(3, cancelled=True)])
        store.save_stocks(con, "wb", [Stock("wb", "SKU1", "Товар 1", "Коледино", 7)])
        with tempfile.TemporaryDirectory() as d:
            p = Path(d) / "s.xlsx"
            report.to_xlsx(con, p)
            wb = load_workbook(p)
            self.assertEqual(wb.sheetnames, ["Заказы", "Остатки", "По дням"])
            days = list(wb["По дням"].iter_rows(min_row=2, values_only=True))
            self.assertEqual(days[0][2:], (1, 1000.0, 1))  # сегодня: 1 заказ, 1 отмена


class NotifyTest(unittest.TestCase):
    def test_pack(self):
        msgs = notify.pack([notify.new_order(order(i)) for i in range(200)])
        self.assertGreater(len(msgs), 1)
        self.assertTrue(all(len(m) <= notify.CHUNK for m in msgs))
        self.assertEqual(sum(m.count("Новый заказ") for m in msgs), 200)

    def test_money(self):
        self.assertEqual(notify.money(Decimal("84300")), "84 300 ₽")


if __name__ == "__main__":
    unittest.main()
